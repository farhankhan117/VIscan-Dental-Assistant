"""Extract searchable text from reports uploaded for the active chatbot session."""
from io import BytesIO
from pathlib import Path
import re
import uuid
import zipfile

from docx import Document
import pymupdf as fitz
from PIL import Image, ImageSequence, UnidentifiedImageError
import pytesseract


_IMAGE_FORMATS = {
    ".apng": {"PNG"},
    ".bmp": {"BMP"},
    ".gif": {"GIF"},
    ".ico": {"ICO"},
    ".jfif": {"JPEG"},
    ".jpe": {"JPEG"},
    ".jpg": {"JPEG"},
    ".jpeg": {"JPEG"},
    ".pbm": {"PPM"},
    ".pcx": {"PCX"},
    ".pfm": {"PPM"},
    ".pgm": {"PPM"},
    ".png": {"PNG"},
    ".pnm": {"PPM"},
    ".ppm": {"PPM"},
    ".tga": {"TGA"},
    ".tif": {"TIFF"},
    ".tiff": {"TIFF"},
    ".webp": {"WEBP"},
}
_MAX_TEXT_CHARS = 100_000
_MAX_PAGES = 100


def _ocr_image(image):
    try:
        return pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError as exc:
        raise ValueError(
            "Image text extraction requires Tesseract OCR. Install Tesseract and restart the API."
        ) from exc
    except pytesseract.TesseractError as exc:
        raise ValueError("The uploaded image could not be read by OCR.") from exc


def _extract_pdf(data):
    try:
        document = fitz.open(stream=data, filetype="pdf")
    except (fitz.FileDataError, RuntimeError) as exc:
        raise ValueError("The uploaded file is not a valid PDF.") from exc
    with document:
        if document.needs_pass:
            raise ValueError("Password-protected PDF reports are not supported.")
        if document.page_count > _MAX_PAGES:
            raise ValueError(f"PDF reports are limited to {_MAX_PAGES} pages.")
        page_text = []
        for page in document:
            text = page.get_text("text").strip()
            if not text:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                image = Image.frombytes(
                    "RGB",
                    (pixmap.width, pixmap.height),
                    pixmap.samples,
                )
                text = _ocr_image(image).strip()
            if text:
                page_text.append(text)
    return "\n".join(page_text)


def _extract_docx(data):
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if (
                len(members) > 2000
                or sum(member.file_size for member in members) > 50 * 1024 * 1024
                or "word/document.xml" not in archive.namelist()
            ):
                raise ValueError("The uploaded file is not a valid or supported DOCX document.")
        document = Document(BytesIO(data))
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("The uploaded file is not a valid DOCX document.") from exc
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _extract_image(data, extension):
    try:
        image = Image.open(BytesIO(data))
        if image.format not in _IMAGE_FORMATS[extension]:
            raise ValueError("The file contents do not match the image extension.")
        if getattr(image, "n_frames", 1) > _MAX_PAGES:
            raise ValueError(f"Image reports are limited to {_MAX_PAGES} frames.")
        if image.width * image.height > 40_000_000:
            raise ValueError("Image dimensions are too large to process safely.")
        text = []
        for frame in ImageSequence.Iterator(image):
            text.append(_ocr_image(frame.convert("RGB")).strip())
        return "\n".join(part for part in text if part)
    except UnidentifiedImageError as exc:
        raise ValueError("The uploaded file is not a supported image.") from exc
    except (OSError, Image.DecompressionBombError) as exc:
        raise ValueError("The uploaded image could not be decoded.") from exc


def extract_report(filename, data):
    """Return a chatbot-compatible report dictionary without persisting the upload."""
    extension = Path(filename).suffix.casefold()
    if extension not in {".pdf", ".docx", *_IMAGE_FORMATS}:
        raise ValueError("Supported report formats are PDF, DOCX, and common image files.")
    if not data:
        raise ValueError("The uploaded report is empty.")

    if extension == ".pdf":
        text = _extract_pdf(data)
        report_type = "Uploaded PDF"
    elif extension == ".docx":
        text = _extract_docx(data)
        report_type = "Uploaded DOCX"
    else:
        text = _extract_image(data, extension)
        report_type = "Uploaded image"

    text = text[:_MAX_TEXT_CHARS].strip()
    if not text:
        raise ValueError(
            "No readable text was found in the upload. Check the document or image quality and try again."
        )

    findings = [
        sentence.strip()
        for sentence in re.split(r"(?<=[.!?])\s+|\n+", text)
        if sentence.strip()
    ]
    display_name = Path(filename.replace("\\", "/")).name
    return {
        "report_id": f"UPLOAD-{uuid.uuid4().hex[:12].upper()}",
        "patient_id": None,
        "patient_name": display_name,
        "report_type": report_type,
        "date": None,
        "findings": findings,
        "report_text": text,
        "synthetic": False,
        "uploaded": True,
    }
